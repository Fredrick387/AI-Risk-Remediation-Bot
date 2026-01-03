import pandas as pd
from openai import OpenAI
import json
import os
from datetime import datetime
from colorama import Fore, init
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from _keys import OPENAI_API_KEY

init(autoreset=True)
client = OpenAI(api_key=OPENAI_API_KEY)

# ----------------------------
# Timestamped output names
# ----------------------------
run_timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
csv_output_file = f"analyzed_scan_{run_timestamp}.csv"
xlsx_output_file = f"analyzed_scan_{run_timestamp}.xlsx"

# ----------------------------
# Risk / Difficulty weights
# ----------------------------
risk_weight = {
    'Critical': 4,
    'High': 3,
    'Medium': 2,
    'Low': 1
}

difficulty_weight = {
    'Easy': 1,
    'Medium': 2,
    'Hard': 3
}

# ----------------------------
# Helper functions
# ----------------------------
def clean_sentence(text, max_len=250):
    if not isinstance(text, str):
        return ''
    text = text.strip()
    return text if len(text) <= max_len else text[:max_len].rsplit(' ', 1)[0] + '...'

def clean_title(text):
    return text.strip().rstrip('.') if isinstance(text, str) else ''

# ----------------------------
# Locate CSV input (SAFE)
# ----------------------------
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))

csv_files = [
    f for f in os.listdir(SCRIPT_DIR)
    if f.lower().endswith('.csv')
    and not f.startswith('analyzed_scan_')
]

if not csv_files:
    print(f"{Fore.RED}No CSV files found in script directory.")
    exit(1)

if len(csv_files) > 1:
    print(f"{Fore.YELLOW}Multiple CSVs found:")
    for i, f in enumerate(csv_files, start=1):
        print(f"  {i}: {f}")
    choice = int(input("Select CSV #: ")) - 1
    filename = csv_files[choice]
else:
    filename = csv_files[0]

filename = os.path.join(SCRIPT_DIR, filename)
print(f"{Fore.CYAN}Analyzing: {os.path.basename(filename)}")

print(f"{Fore.CYAN}Analyzing: {filename}")

df = pd.read_csv(filename)

if df.empty:
    print(f"{Fore.RED}CSV loaded but contains no rows.")
    exit(1)


# ----------------------------
# Filter compliance data
# ----------------------------
compliance = df[df['Plugin Family'] == 'Policy Compliance'].copy()

compliance['Audit ID'] = (
    compliance['Name'].str.extract(r'(\d+(?:\.\d+)+)')[0]
    .combine_first(compliance['Synopsis'].str.extract(r'(\d+(?:\.\d+)+)')[0])
    .combine_first(compliance['Description'].str.extract(r'(\d+(?:\.\d+)+)')[0])
)

compliance = compliance[compliance['Audit ID'].notna()]

print(f"{Fore.YELLOW}Valid compliance audits: {len(compliance)}")

# LIMIT FOR TESTING
to_analyze = compliance.head(20)

results = []

# ----------------------------
# AI analysis loop
# ----------------------------
for _, row in to_analyze.iterrows():
    audit_id = row['Audit ID']
    name = row['Name']

    print(f"{Fore.WHITE}Analyzing {audit_id}")

    prompt = f"""
You are analyzing a CIS or STIG compliance control.

Audit ID: {audit_id}
Scanner Name: {name}

Synopsis:
{row.get('Synopsis', '')}

Description:
{row.get('Description', '')}

Solution:
{row.get('Solution', '')}

Return strict JSON only:
{{
  "audit_title": "",
  "brief_description": "",
  "proposed_solution": "",
  "risk": "Low/Medium/High/Critical",
  "difficulty": "Easy/Medium/Hard",
  "quick_win": "Yes/No",
  "rationale": ""
}}
"""

    resp = client.chat.completions.create(
        model="gpt-4o-mini",
        messages=[{"role": "user", "content": prompt}]
    )

    try:
        content = resp.choices[0].message.content
        analysis = json.loads(content[content.find('{'):content.rfind('}') + 1])

        priority_score = (
            risk_weight.get(analysis.get('risk'), 0) * 10
            - difficulty_weight.get(analysis.get('difficulty'), 0)
        )

        results.append({
            'Audit ID': audit_id,
            'Audit Title': clean_title(analysis.get('audit_title')),
            'Risk': analysis.get('risk'),
            'Difficulty': analysis.get('difficulty'),
            'Quick Win': analysis.get('quick_win'),
            'Priority Score': priority_score,
            'Brief Description': clean_sentence(analysis.get('brief_description')),
            'Proposed Solution': clean_sentence(analysis.get('proposed_solution')),
            'Rationale': clean_sentence(analysis.get('rationale'))
        })

        print(f"{Fore.GREEN}✓ {analysis.get('risk')} | Priority {priority_score}")

    except Exception as e:
        print(f"{Fore.RED}Parse error: {e}")

# ----------------------------
# Create DataFrame
# ----------------------------
columns = [
    'Audit ID',
    'Audit Title',
    'Risk',
    'Difficulty',
    'Quick Win',
    'Priority Score',
    'Brief Description',
    'Proposed Solution',
    'Rationale'
]

output_df = pd.DataFrame(results)[columns]
output_df.to_csv(csv_output_file, index=False)
output_df.to_excel(xlsx_output_file, index=False)

# ----------------------------
# Excel formatting
# ----------------------------
wb = load_workbook(xlsx_output_file)
ws = wb.active
from openpyxl.styles import Alignment

ws.freeze_panes = 'A2'
ws.auto_filter.ref = ws.dimensions

header_fill = PatternFill(start_color='DDDDDD', end_color='DDDDDD', fill_type='solid')
header_font = Font(bold=True)

for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center')

wrap_columns = {
    'Brief Description': 50,
    'Proposed Solution': 50,
    'Rationale': 45
}

for col_idx, col_name in enumerate(output_df.columns, start=1):
    col_letter = get_column_letter(col_idx)

    if col_name in wrap_columns:
        ws.column_dimensions[col_letter].width = wrap_columns[col_name]
        for row in range(2, ws.max_row + 1):
            ws.cell(row=row, column=col_idx).alignment = Alignment(
                wrap_text=True,
                vertical='top'
            )
    else:
        ws.column_dimensions[col_letter].width = 22

# Risk color mapping
risk_colors = {
    'Critical': 'FF4C4C',
    'High': 'FFA500',
    'Medium': 'FFD966',
    'Low': '92D050'
}

risk_col = output_df.columns.get_loc('Risk') + 1
for row in range(2, ws.max_row + 1):
    cell = ws.cell(row=row, column=risk_col)
    if cell.value in risk_colors:
        cell.fill = PatternFill(
            start_color=risk_colors[cell.value],
            end_color=risk_colors[cell.value],
            fill_type='solid'
        )

# Wrap text + size Audit Title column
audit_title_col = output_df.columns.get_loc('Audit Title') + 1
audit_title_letter = get_column_letter(audit_title_col)

# Set column width (tuned for 1–2 line wrap)
ws.column_dimensions[audit_title_letter].width = 40

# Apply wrap + top alignment
for row in range(2, ws.max_row + 1):
    ws.cell(row=row, column=audit_title_col).alignment = Alignment(
        wrap_text=True,
        vertical='top'
    )


wb.save(xlsx_output_file)

print(f"{Fore.GREEN}Reports created:")
print(f"  CSV  → {csv_output_file}")
print(f"  XLSX → {xlsx_output_file}")
